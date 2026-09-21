"""Kiểm thử luồng nhận CV.

Hai thứ được soi kỹ nhất, vì hỏng ở đó thì hỏng âm thầm:

1. **Đoạn dẫn phải có thật trong CV.** Đây là ranh giới giữa "máy đọc được" và
   "máy đoán ra". Bỏ chốt này thì hồ sơ vẫn đầy đủ trường, trông vẫn đẹp, nhưng
   có thể toàn số liệu bịa — và không ai phát hiện ra cho tới lúc gọi điện cho
   ứng viên.
2. **Hỏng ở bước bóc tách vẫn phải còn dấu vết.** Ứng viên đã gửi file thì hệ
   thống phải nhớ là có file, kể cả khi không đọc nổi.
"""
import io
import unittest
from unittest.mock import AsyncMock, patch

from app.documents import extractor, reader, storage
from app.services import cv_service


SESSION = "phien-ung-vien-0001"

CV_TEXT = """CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM
SƠ YẾU LÝ LỊCH
Họ và tên: Nguyễn Thị Lan
Năm sinh: 1999
Giới tính: Nữ
Trình độ: Cao đẳng Điều dưỡng, trường Cao đẳng Y tế Hà Nội
Tiếng Nhật: đã có chứng chỉ N4
Kinh nghiệm: 3 năm chăm sóc người cao tuổi tại viện dưỡng lão Thanh Xuân
Điện thoại: 0912345678
"""


def make_pdf(text: str) -> bytes:
    import pymupdf

    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((50, 60), text, fontsize=9, fontname="helv")
    data = document.tobytes()
    document.close()
    return data


def make_docx(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    from docx import Document

    document = Document()
    for line in paragraphs:
        document.add_paragraph(line)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_index, row in enumerate(table_rows):
            for cell_index, value in enumerate(row):
                table.cell(row_index, cell_index).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class ReaderTests(unittest.TestCase):
    def test_nhan_dien_theo_noi_dung_khong_theo_duoi_file(self):
        """File thi hành đổi tên thành .pdf vẫn bị chặn."""
        with self.assertRaises(reader.UnsupportedDocument):
            reader.detect_kind(b"MZ\x90\x00 day la file exe", "application/pdf", "cv.pdf")

    def test_pdf_that_duoc_nhan(self):
        data = make_pdf("Ho va ten: Nguyen Thi Lan")
        self.assertEqual(
            reader.detect_kind(data, "application/octet-stream", "cv.pdf"),
            reader.PDF_MIME,
        )

    def test_doc_duoc_chu_trong_pdf(self):
        # Chỉ ASCII: font mặc định của PDF không dựng được dấu tiếng Việt, mà bài
        # kiểm thử này hỏi về việc rút chữ chứ không hỏi về phông chữ.
        kind, result = reader.read(make_pdf("Ho va ten: Nguyen Thi Lan" + " x" * 60))
        self.assertEqual(kind, reader.PDF_MIME)
        self.assertIn("Nguyen Thi Lan", result.text)
        self.assertEqual(result.page_count, 1)
        self.assertFalse(result.needs_model_ocr)

    def test_pdf_scan_bi_danh_dau_can_mo_hinh_doc(self):
        """PDF không có chữ — đúng hình dạng của một bản scan."""
        _, result = reader.read(make_pdf(""))
        self.assertTrue(result.needs_model_ocr)
        self.assertEqual(result.text, "")

    def test_nhan_dien_cac_dinh_dang_anh(self):
        """Nhận theo byte đầu, không theo đuôi — kể cả HEIC của iPhone."""
        cases = [
            (bytes.fromhex("ffd8ffe0"), reader.JPEG_MIME),
            (bytes.fromhex("89504e470d0a1a0a"), reader.PNG_MIME),
            (b"RIFF" + bytes(4) + b"WEBP", reader.WEBP_MIME),
            (bytes.fromhex("00000018") + b"ftypheic", reader.HEIC_MIME),
            (bytes.fromhex("00000018") + b"ftypmif1", reader.HEIC_MIME),
        ]
        for prefix, expected in cases:
            with self.subTest(expected=expected):
                data = prefix + b"0" * 40
                self.assertEqual(reader.detect_kind(data, "", "anh"), expected)

    def test_anh_luon_can_mo_hinh_doc(self):
        kind, result = reader.read(bytes.fromhex("ffd8ffe0") + b"0" * 40, filename="cv.jpg")
        self.assertEqual(kind, reader.JPEG_MIME)
        self.assertTrue(result.needs_model_ocr)
        self.assertEqual(result.text, "")

    def test_van_chan_file_khong_phai_tai_lieu_hay_anh(self):
        """Khai là ảnh nhưng nội dung là file thi hành — vẫn chặn."""
        with self.assertRaises(reader.UnsupportedDocument):
            reader.detect_kind(b"MZ" + bytes(2) + b" exe", "image/jpeg", "cv.jpg")

    def test_docx_lay_ca_noi_dung_trong_bang(self):
        data = make_docx(
            ["SƠ YẾU LÝ LỊCH", "Họ và tên: Nguyễn Thị Lan"],
            table_rows=[["Trình độ tiếng Nhật", "N4"], ["Kinh nghiệm", "3 năm"]],
        )
        kind, result = reader.read(data, filename="cv.docx")
        self.assertEqual(kind, reader.DOCX_MIME)
        self.assertIn("Nguyễn Thị Lan", result.text)
        self.assertIn("Trình độ tiếng Nhật | N4", result.text)

    def test_cat_bot_van_ban_qua_dai(self):
        long_text = "a" * (reader.MAX_TEXT_CHARS + 5_000)
        self.assertEqual(len(reader._clean(long_text)), reader.MAX_TEXT_CHARS)


class MapExtractionTests(unittest.TestCase):
    """Phần lọc kết quả của mô hình. Thuần — không cần mạng, không cần khóa API."""

    def test_nhan_truong_co_doan_dan_that(self):
        raw = {
            "full_name": "Nguyễn Thị Lan",
            "japanese_level": "N4",
            "evidence": {
                "full_name": "Họ và tên: Nguyễn Thị Lan",
                "japanese_level": "Tiếng Nhật: đã có chứng chỉ N4",
            },
        }
        result = extractor.map_extraction(raw, CV_TEXT)
        self.assertEqual(result.fields["full_name"], "Nguyễn Thị Lan")
        self.assertEqual(result.fields["japanese_level"], "N4")
        self.assertEqual(result.rejected, {})

    def test_loai_truong_co_doan_dan_khong_ton_tai(self):
        """Trường hợp đáng sợ nhất: giá trị nghe rất hợp lý nhưng CV không hề nói."""
        raw = {
            "birth_year": 1995,
            "evidence": {"birth_year": "Năm sinh: 1995"},
        }
        result = extractor.map_extraction(raw, CV_TEXT)
        self.assertNotIn("birth_year", result.fields)
        self.assertEqual(result.rejected["birth_year"], "đoạn dẫn không có trong CV")

    def test_loai_truong_khong_co_doan_dan(self):
        result = extractor.map_extraction({"full_name": "Nguyễn Thị Lan"}, CV_TEXT)
        self.assertEqual(result.fields, {})
        self.assertEqual(result.rejected["full_name"], "không chỉ được đoạn dẫn trong CV")

    def test_doan_dan_khac_khoang_trang_van_duoc_nhan(self):
        raw = {
            "full_name": "Nguyễn Thị Lan",
            "evidence": {"full_name": "họ và tên:    Nguyễn Thị Lan"},
        }
        result = extractor.map_extraction(raw, CV_TEXT)
        self.assertIn("full_name", result.fields)

    def test_loai_gia_tri_ngoai_danh_muc(self):
        raw = {
            "japanese_level": "N9",
            "evidence": {"japanese_level": "Tiếng Nhật: đã có chứng chỉ N4"},
        }
        result = extractor.map_extraction(raw, CV_TEXT)
        self.assertEqual(result.fields, {})
        self.assertIn("ngoài danh mục", result.rejected["japanese_level"])

    def test_loai_truong_khong_thuoc_ho_so(self):
        raw = {"luong_mong_muon": 200000, "evidence": {"luong_mong_muon": "abc"}}
        result = extractor.map_extraction(raw, CV_TEXT)
        self.assertEqual(result.rejected["luong_mong_muon"], "không phải trường của hồ sơ")

    def test_chuan_hoa_gia_tri_ve_ma_danh_muc(self):
        raw = {
            "gender": "Nữ",
            "education_level": "Cao đẳng",
            "phone": "0912345678",
            "evidence": {
                "gender": "Giới tính: Nữ",
                "education_level": "Trình độ: Cao đẳng Điều dưỡng",
                "phone": "Điện thoại: 0912345678",
            },
        }
        result = extractor.map_extraction(raw, CV_TEXT)
        self.assertEqual(result.fields["gender"], "nu")
        self.assertEqual(result.fields["education_level"], "cao_dang")
        self.assertEqual(result.fields["phone"], "0912345678")

    def test_nhan_doan_dan_dang_mang(self):
        """Hình dạng lược đồ đang dùng thật: mảng `[{field, quote}]`."""
        raw = {
            "full_name": "Nguyễn Thị Lan",
            "evidence": [{"field": "full_name", "quote": "Họ và tên: Nguyễn Thị Lan"}],
        }
        result = extractor.map_extraction(raw, CV_TEXT)
        self.assertEqual(result.fields["full_name"], "Nguyễn Thị Lan")

    def test_doan_dan_dang_mang_van_bi_soi(self):
        raw = {
            "birth_year": 1995,
            "evidence": [{"field": "birth_year", "quote": "Năm sinh: 1995"}],
        }
        result = extractor.map_extraction(raw, CV_TEXT)
        self.assertEqual(result.rejected["birth_year"], "đoạn dẫn không có trong CV")

    def test_nam_sinh_ngoai_khoang_bi_loai(self):
        raw = {"birth_year": 1899, "evidence": {"birth_year": "SƠ YẾU LÝ LỊCH"}}
        result = extractor.map_extraction(raw, CV_TEXT)
        self.assertEqual(result.fields, {})
        self.assertIn("birth_year", result.rejected)


class StorageTests(unittest.TestCase):
    def test_khong_ghi_de_file_da_ton_tai(self):
        import tempfile
        from pathlib import Path

        with tempfile.TemporaryDirectory() as folder:
            with patch.object(
                type(storage.settings), "cv_storage_path", property(lambda _: Path(folder))
            ):
                storage.save(b"lan mot", "CV-AAAAAA", ".pdf")
                with self.assertRaises(FileExistsError):
                    storage.save(b"lan hai", "CV-AAAAAA", ".pdf")

    def test_van_tay_doi_khi_noi_dung_doi(self):
        self.assertNotEqual(storage.digest(b"mot"), storage.digest(b"hai"))
        self.assertEqual(storage.digest(b"mot"), storage.digest(b"mot"))


class IngestTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.extraction = extractor.Extraction(
            fields={"full_name": "Nguyễn Thị Lan", "japanese_level": "N4"},
            evidence={
                "full_name": "Họ và tên: Nguyễn Thị Lan",
                "japanese_level": "Tiếng Nhật: đã có chứng chỉ N4",
            },
        )

    def _patches(self, **overrides):
        """Khung mock chung: đĩa, database và mô hình đều thay bằng bản giả."""
        defaults = {
            "count_in_session": AsyncMock(return_value=0),
            "find_in_session": AsyncMock(return_value=None),
            "create": AsyncMock(side_effect=lambda doc: dict(doc)),
            "mark_extracted": AsyncMock(),
            "mark_failed": AsyncMock(),
            "get_by_session": AsyncMock(return_value=None),
            "create_profile": AsyncMock(
                side_effect=lambda doc: {**doc, "version": 1, "history": []}
            ),
            "attach_document": AsyncMock(),
            "apply_changes": AsyncMock(),
            "save": lambda data, code, ext: f"2026/09/{code}{ext}",
            "extract_fields": lambda text: self.extraction,
        }
        defaults.update(overrides)
        return defaults

    def _apply(self, mocks):
        return (
            patch.object(cv_service.documents, "count_in_session", mocks["count_in_session"]),
            patch.object(cv_service.documents, "find_in_session", mocks["find_in_session"]),
            patch.object(cv_service.documents, "create", mocks["create"]),
            patch.object(cv_service.documents, "mark_extracted", mocks["mark_extracted"]),
            patch.object(cv_service.documents, "mark_failed", mocks["mark_failed"]),
            patch.object(cv_service.profiles, "get_by_session", mocks["get_by_session"]),
            patch.object(cv_service.profiles, "create_profile", mocks["create_profile"]),
            patch.object(cv_service.profiles, "attach_document", mocks["attach_document"]),
            patch.object(cv_service.profiles, "apply_changes", mocks["apply_changes"]),
            patch.object(cv_service.storage, "save", mocks["save"]),
            patch.object(cv_service.extractor, "extract_fields", mocks["extract_fields"]),
        )

    async def _ingest(self, data: bytes, mocks: dict, filename="cv.docx"):
        from contextlib import ExitStack

        with ExitStack() as stack:
            for item in self._apply(mocks):
                stack.enter_context(item)
            return await cv_service.ingest(
                data=data, filename=filename, content_type="", session_id=SESSION
            )

    async def test_tu_choi_file_qua_lon(self):
        oversized = b"%PDF-" + b"0" * (cv_service.settings.max_upload_bytes + 1)
        with self.assertRaises(cv_service.UploadRejected) as ctx:
            await cv_service.ingest(
                data=oversized, filename="cv.pdf", content_type="", session_id=SESSION
            )
        self.assertIn("MB", str(ctx.exception))

    async def test_tu_choi_dinh_dang_la(self):
        with patch.object(cv_service.documents, "count_in_session", AsyncMock(return_value=0)):
            with self.assertRaises(cv_service.UploadRejected):
                await cv_service.ingest(
                    data=b"anh jpeg gia dinh",
                    filename="cv.jpg",
                    content_type="image/jpeg",
                    session_id=SESSION,
                )

    async def test_tu_choi_khi_vuot_han_muc_moi_phien(self):
        mocks = self._patches(
            count_in_session=AsyncMock(return_value=cv_service.documents.MAX_PER_SESSION)
        )
        with patch.object(cv_service.documents, "count_in_session", mocks["count_in_session"]):
            with self.assertRaises(cv_service.UploadRejected):
                await cv_service.ingest(
                    data=make_docx(["x"]), filename="cv.docx", content_type="", session_id=SESSION
                )

    async def test_tao_ho_so_cho_so_moi_o_trang_thai_cho_xac_nhan(self):
        mocks = self._patches()
        data = make_docx([CV_TEXT])
        result = await self._ingest(data, mocks)

        created = mocks["create_profile"].await_args.args[0]
        self.assertEqual(created["status"], cv_service.profiles.STATUS_EXTRACTED)
        self.assertEqual(created["fields"]["full_name"]["source"], "cv")
        self.assertEqual(
            created["fields"]["full_name"]["evidence"], "Họ và tên: Nguyễn Thị Lan"
        )
        self.assertEqual(result.accepted_fields, ["full_name", "japanese_level"])
        mocks["attach_document"].assert_awaited_once()

    async def test_gui_lai_dung_file_cu_khong_goi_lai_mo_hinh(self):
        called = []

        def spy(text):
            called.append(text)
            return self.extraction

        mocks = self._patches(
            find_in_session=AsyncMock(
                return_value={
                    "code": "CV-OLD123",
                    "status": cv_service.documents.STATUS_EXTRACTED,
                    "extracted_fields": ["full_name"],
                    "rejected": {},
                    "stored_path": "2026/09/CV-OLD123.docx",
                }
            ),
            extract_fields=spy,
        )
        result = await self._ingest(make_docx([CV_TEXT]), mocks)

        self.assertEqual(called, [])
        self.assertEqual(result.document["code"], "CV-OLD123")
        self.assertIn("đã được tiếp nhận", result.message)
        mocks["create"].assert_not_awaited()

    async def test_gui_lai_file_tung_doc_hong_thi_doc_lai(self):
        """Lần trước hỏng vì trục trặc nhất thời — gửi lại phải được thử lại."""
        mocks = self._patches(
            find_in_session=AsyncMock(
                return_value={
                    "code": "CV-OLD999",
                    "status": cv_service.documents.STATUS_FAILED,
                    "extracted_fields": [],
                    "rejected": {},
                }
            )
        )
        result = await self._ingest(make_docx([CV_TEXT]), mocks)

        mocks["create"].assert_awaited_once()
        self.assertEqual(result.accepted_fields, ["full_name", "japanese_level"])

    async def test_file_scan_van_duoc_ghi_nhan(self):
        mocks = self._patches()
        result = await self._ingest(make_pdf(""), mocks, filename="cv-scan.pdf")

        mocks["create"].assert_awaited_once()
        code, status, _ = _mark_failed_args(mocks["mark_failed"])
        self.assertEqual(status, cv_service.documents.STATUS_UNREADABLE)
        self.assertTrue(code.startswith("CV-"))
        self.assertEqual(result.accepted_fields, [])

    async def test_nhan_anh_chup_va_ghi_nhan_de_do_nhu_cau(self):
        """Ảnh được cất giữ, không gọi mô hình, và ứng viên biết file đã tới nơi."""
        called = []

        def spy(text):
            called.append(text)
            return self.extraction

        mocks = self._patches(extract_fields=spy)
        data = bytes.fromhex("ffd8ffe0") + b"0" * 2000
        result = await self._ingest(data, mocks, filename="cv.jpg")

        self.assertEqual(called, [])
        mocks["create"].assert_awaited_once()
        self.assertEqual(
            mocks["create"].await_args.args[0]["content_type"], reader.JPEG_MIME
        )
        _, status, _ = _mark_failed_args(mocks["mark_failed"])
        self.assertEqual(status, cv_service.documents.STATUS_UNREADABLE)
        self.assertIn("đã nhận ảnh", result.message)

    async def test_boc_tach_hong_van_giu_ban_ghi(self):
        def fail(text):
            raise extractor.ExtractionFailed("het han muc")

        mocks = self._patches(extract_fields=fail)
        result = await self._ingest(make_docx([CV_TEXT]), mocks)

        mocks["create"].assert_awaited_once()
        _, status, error = _mark_failed_args(mocks["mark_failed"])
        self.assertEqual(status, cv_service.documents.STATUS_FAILED)
        self.assertIn("het han muc", error)
        self.assertIsNone(result.document.get("profile_code"))

    async def test_khong_de_len_gia_tri_ung_vien_da_tu_sua(self):
        """Ưu tiên nguồn: `user_confirmed` mạnh hơn `cv`."""
        existing = {
            "code": "UV-ABC123",
            "session_id": SESSION,
            "version": 2,
            "fields": {
                "full_name": cv_service.profiles.cell("Nguyễn Thị Lan Anh", "user_confirmed")
            },
            "preferences": {},
            "history": [],
            "status": cv_service.profiles.STATUS_EXTRACTED,
        }
        mocks = self._patches(
            get_by_session=AsyncMock(return_value=existing),
            apply_changes=AsyncMock(side_effect=lambda session_id, **kw: {**existing, **kw}),
        )
        await self._ingest(make_docx([CV_TEXT]), mocks)

        merged = mocks["apply_changes"].await_args.kwargs["fields"]
        self.assertEqual(merged["full_name"]["value"], "Nguyễn Thị Lan Anh")
        self.assertEqual(merged["full_name"]["source"], "user_confirmed")
        # Trường ứng viên chưa khai thì bản đọc từ CV vẫn được điền vào.
        self.assertEqual(merged["japanese_level"]["value"], "N4")


def _mark_failed_args(mock) -> tuple[str, str, str]:
    args = mock.await_args.args
    return args[0], args[1], args[2]


if __name__ == "__main__":
    unittest.main()
