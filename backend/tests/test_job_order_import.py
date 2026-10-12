"""Kiểm thử nhập danh mục đơn tuyển dụng từ Excel.

File Excel là dữ liệu do người gõ tay, nên phần lớn ca kiểm thử ở đây nói về
chuyện sai: sai chính tả danh mục, ngày viết nhiều kiểu, ô ghi hai mức tiếng
Nhật, mã đơn lặp. Yêu cầu xuyên suốt là **báo đúng dòng và đúng lý do**, vì người
sửa file là nhân viên nghiệp vụ chứ không phải lập trình viên.
"""
import unittest
from datetime import date, datetime, timedelta
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.core.timeutil import local_today
from app.services import job_order_import as importer


def workbook_bytes(rows: list[dict], headers: list[str] | None = None) -> bytes:
    """Dựng một file Excel tối giản đúng định dạng bộ nhập mong đợi."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = importer.DATA_SHEET
    labels = headers or [column.label for column in importer.COLUMNS]
    sheet.append(labels)
    key_by_label = {column.label: column.key for column in importer.COLUMNS}
    for row in rows:
        sheet.append([row.get(key_by_label.get(label, label)) for label in labels])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def valid_row(**overrides) -> dict:
    row = {
        "title": "Điều dưỡng viện dưỡng lão Tokyo",
        "employer_name": "Viện dưỡng lão Sakura",
        "employer_type": "Viện dưỡng lão",
        "program": "Tokutei Ginou",
        "prefecture": "Tokyo",
        "quota": 5,
        "japanese_required": "N4",
        "education_required": "Cao đẳng",
        "experience_min": 0,
        "age_min": 20,
        "age_max": 35,
        "gender_pref": "Không yêu cầu",
        "deadline": (local_today() + timedelta(days=60)).strftime("%d/%m/%Y"),
        "salary_min": 190000,
        "salary_max": 210000,
        "allowances": "Ký túc xá; Phụ cấp ca đêm",
        "status": "Đang tuyển",
        "published": "Có",
    }
    row.update(overrides)
    return row


class TemplateTests(unittest.TestCase):
    def test_template_contains_every_column_and_the_guide(self):
        workbook = load_workbook(BytesIO(importer.template_bytes()))
        self.assertIn(importer.DATA_SHEET, workbook.sheetnames)
        self.assertIn(importer.GUIDE_SHEET, workbook.sheetnames)
        self.assertIn(importer.CATALOG_SHEET, workbook.sheetnames)
        header = [cell.value for cell in workbook[importer.DATA_SHEET][1]]
        self.assertEqual(header, [column.label for column in importer.COLUMNS])

    def test_sample_sheet_rows_survive_a_round_trip(self):
        """Dữ liệu ví dụ trong file mẫu phải nhập lại được, nếu không thì mẫu sai."""
        workbook = load_workbook(BytesIO(importer.template_bytes()))
        samples = workbook[importer.SAMPLE_SHEET]
        rebuilt = Workbook()
        sheet = rebuilt.active
        sheet.title = importer.DATA_SHEET
        for row in samples.iter_rows(values_only=True):
            sheet.append(list(row))
        buffer = BytesIO()
        rebuilt.save(buffer)

        preview = importer.parse_workbook(buffer.getvalue())
        self.assertEqual(preview.summary["error"], 0, msg=str(
            [row.errors for row in preview.rows if row.errors]
        ))
        self.assertGreaterEqual(preview.summary["total"], 10)


class ImportHappyPathTests(unittest.TestCase):
    def test_valid_row_is_parsed_into_a_document(self):
        preview = importer.parse_workbook(workbook_bytes([valid_row()]))
        self.assertEqual(preview.summary, {"total": 1, "create": 1, "update": 0, "error": 0})
        data = preview.rows[0].data
        self.assertEqual(data["employer_type"], "vien_duong_lao")
        self.assertEqual(data["program"], "tokutei_ginou")
        self.assertEqual(data["region_group"], "kanto")
        self.assertEqual(data["requirements"]["japanese_required"], "N4")
        self.assertEqual(data["reference"]["allowances"], ["Ký túc xá", "Phụ cấp ca đêm"])
        self.assertTrue(data["published"])

    def test_labels_without_diacritics_are_accepted(self):
        row = valid_row(employer_type="vien duong lao", gender_pref="khong yeu cau", published="co")
        preview = importer.parse_workbook(workbook_bytes([row]))
        self.assertEqual(preview.rows[0].action, "create")
        self.assertEqual(preview.rows[0].data["employer_type"], "vien_duong_lao")

    def test_a_row_with_a_code_is_treated_as_an_update(self):
        preview = importer.parse_workbook(workbook_bytes([valid_row(code="DH-0007")]))
        self.assertEqual(preview.rows[0].action, "update")
        self.assertEqual(preview.rows[0].code, "DH-0007")

    def test_several_date_formats_are_understood(self):
        future = local_today() + timedelta(days=45)
        for value in (
            future.strftime("%d/%m/%Y"),
            future.isoformat(),
            future.strftime("%d-%m-%Y"),
            datetime(future.year, future.month, future.day),
        ):
            preview = importer.parse_workbook(workbook_bytes([valid_row(deadline=value)]))
            self.assertEqual(
                preview.rows[0].action, "create", msg=f"{value!r}: {preview.rows[0].errors}"
            )
            self.assertEqual(preview.rows[0].data["deadline"], future.isoformat())

    def test_blank_rows_are_skipped_not_reported_as_errors(self):
        preview = importer.parse_workbook(workbook_bytes([valid_row(), {}, valid_row()]))
        self.assertEqual(preview.summary["total"], 2)

    def test_column_order_does_not_matter(self):
        labels = [column.label for column in importer.COLUMNS]
        reversed_labels = list(reversed(labels))
        preview = importer.parse_workbook(workbook_bytes([valid_row()], headers=reversed_labels))
        self.assertEqual(preview.rows[0].action, "create")


class ImportErrorReportingTests(unittest.TestCase):
    def _errors(self, **overrides) -> list[str]:
        preview = importer.parse_workbook(workbook_bytes([valid_row(**overrides)]))
        self.assertEqual(preview.rows[0].action, "error")
        return preview.rows[0].errors

    def test_missing_required_column_is_reported_before_any_row(self):
        labels = [
            column.label for column in importer.COLUMNS if column.key != "title"
        ]
        preview = importer.parse_workbook(workbook_bytes([valid_row()], headers=labels))
        self.assertIn("Tên đơn", preview.missing_columns)
        self.assertEqual(preview.rows, [])

    def test_missing_required_value_names_the_column(self):
        errors = self._errors(title=None)
        self.assertTrue(any("tên đơn" in error.lower() for error in errors))

    def test_unknown_catalog_value_lists_the_accepted_choices(self):
        errors = self._errors(employer_type="nhà hàng")
        self.assertTrue(any("Viện dưỡng lão" in error for error in errors))

    def test_ambiguous_japanese_requirement_is_refused_not_guessed(self):
        """Đây là điều kiện loại ứng viên; đoán sai một bậc là loại oan cả nhóm hồ sơ."""
        errors = self._errors(japanese_required="N4 trở lên, ưu tiên N3")
        self.assertTrue(any("nhiều mức" in error for error in errors))

    def test_inverted_age_range_is_refused(self):
        errors = self._errors(age_min=40, age_max=25)
        self.assertTrue(any("Tuổi từ" in error for error in errors))

    def test_inverted_salary_range_is_refused(self):
        errors = self._errors(salary_min=300000, salary_max=100000)
        self.assertTrue(any("Lương từ" in error for error in errors))

    def test_past_deadline_on_an_open_order_is_refused(self):
        errors = self._errors(
            deadline=(local_today() - timedelta(days=5)).strftime("%d/%m/%Y")
        )
        self.assertTrue(any("đã qua" in error for error in errors))

    def test_past_deadline_is_allowed_on_a_draft(self):
        row = valid_row(
            deadline=(local_today() - timedelta(days=5)).strftime("%d/%m/%Y"),
            status="Nháp",
        )
        preview = importer.parse_workbook(workbook_bytes([row]))
        self.assertEqual(preview.rows[0].action, "create")

    def test_unreadable_date_is_refused(self):
        errors = self._errors(deadline="tháng sau")
        self.assertTrue(any("không đọc được thành ngày" in error for error in errors))

    def test_bad_optional_numbers_are_reported_not_silently_dropped(self):
        """Ô sai ở cột không bắt buộc vẫn phải báo, đừng âm thầm biến thành rỗng."""
        errors = self._errors(experience_min="một năm")
        self.assertTrue(any("Kinh nghiệm tối thiểu" in error for error in errors))

        errors = self._errors(cost_total_vnd="khoảng một trăm triệu")
        self.assertTrue(any("Tổng chi phí" in error for error in errors))

        errors = self._errors(interview_date="cuối tháng")
        self.assertTrue(any("Ngày phỏng vấn dự kiến" in error for error in errors))

    def test_closed_statuses_must_go_through_the_admin_screen(self):
        """Đổi sang Đã đóng qua file thì mất lịch sử ai đóng và vì sao."""
        errors = self._errors(status="Đã đóng")
        self.assertTrue(any("màn hình quản trị" in error for error in errors))

    def test_quota_must_be_positive(self):
        errors = self._errors(quota=0)
        self.assertTrue(any("Số lượng" in error for error in errors))

    def test_duplicate_code_inside_the_file_points_at_the_first_row(self):
        rows = [valid_row(code="DH-0009"), valid_row(code="DH-0009")]
        preview = importer.parse_workbook(workbook_bytes(rows))
        self.assertEqual(preview.rows[1].action, "error")
        self.assertTrue(any("dòng 2" in error for error in preview.rows[1].errors))

    def test_row_numbers_match_the_spreadsheet(self):
        """Nhân viên sửa file theo số dòng Excel, nên số phải khớp tuyệt đối."""
        preview = importer.parse_workbook(
            workbook_bytes([valid_row(), valid_row(title=None)])
        )
        self.assertEqual(preview.rows[0].row_number, 2)
        self.assertEqual(preview.rows[1].row_number, 3)

    def test_corrupt_file_gives_a_readable_message(self):
        with self.assertRaises(ValueError) as raised:
            importer.parse_workbook(b"day khong phai file excel")
        self.assertIn("Không đọc được file Excel", str(raised.exception))


class ConsultantNoteColumnTests(unittest.TestCase):
    """Cột "Thông tin cho tư vấn viên" (09/10) và chốt không xóa ghi chú khi file thiếu cột.

    Hai trường ghi chú được phép xóa về rỗng, nên một dòng cập nhật mang `None` là xóa
    thật. File làm trước khi có cột này mà đem nhập để cập nhật đơn thì ghi chú tư vấn
    viên đã nhập trên màn quản trị sẽ mất sạch — ca dưới ghim chuyện đó không xảy ra.
    """

    def test_template_has_the_column(self):
        self.assertIn("Thông tin cho tư vấn viên", [c.label for c in importer.COLUMNS])

    def test_value_is_read_when_the_column_exists(self):
        preview = importer.parse_workbook(workbook_bytes([valid_row(consultant_note="Hỏi chứng chỉ gốc")]))
        self.assertEqual(preview.rows[0].data["consultant_note"], "Hỏi chứng chỉ gốc")

    def test_old_file_without_the_column_does_not_touch_saved_notes(self):
        cu = [c.label for c in importer.COLUMNS if c.key not in ("consultant_note", "internal_note")]
        preview = importer.parse_workbook(workbook_bytes([valid_row(code="DH-0001")], headers=cu))
        row = preview.rows[0]
        self.assertEqual(row.action, "update")
        self.assertNotIn("consultant_note", row.data)
        self.assertNotIn("internal_note", row.data)

    def test_column_present_but_blank_clears_on_purpose(self):
        preview = importer.parse_workbook(workbook_bytes([valid_row(code="DH-0001", consultant_note=None)]))
        self.assertIn("consultant_note", preview.rows[0].data)
        self.assertIsNone(preview.rows[0].data["consultant_note"])

    def test_why_it_matters_none_would_really_be_written(self):
        # Tầng dữ liệu GIỮ `None` của trường ghi chú — tức một khóa `None` lọt xuống là xóa.
        from app.db.common import flatten_update, keep_allowed_nulls
        from app.db.job_orders import NULLABLE_FIELDS

        flat = keep_allowed_nulls(flatten_update({"consultant_note": None}), NULLABLE_FIELDS)
        self.assertEqual(flat, {"consultant_note": None})


class CotVangKhiCapNhatTests(unittest.TestCase):
    """Cập nhật bằng file bị xóa bớt cột không được xóa dữ liệu đang lưu (12/10).

    Đo thử trước khi sửa: file chỉ có các cột bắt buộc, cập nhật DH-0001 → ghi đè rỗng 13
    trường (thành phố, bằng cấp, tuổi, lương, phụ cấp, chi phí, ngày phỏng vấn, xuất cảnh,
    điểm nổi bật, mô tả), đặt kinh nghiệm về 0 và **tắt công khai**.
    """

    BAT_BUOC = [c.label for c in importer.COLUMNS if c.required] + ["Mã đơn"]

    def test_cap_nhat_chi_ghi_dung_cot_co_trong_file(self):
        from app.db.common import flatten_update, keep_allowed_nulls
        from app.db.job_orders import NULLABLE_FIELDS

        row = importer.parse_workbook(workbook_bytes([valid_row(code="DH-0001")], headers=self.BAT_BUOC)).rows[0]
        self.assertEqual(row.action, "update")
        se_ghi = set(keep_allowed_nulls(flatten_update(row.data), NULLABLE_FIELDS))
        for khong_duoc in ("published", "status", "description", "city", "reference.salary_min",
                           "reference.allowances", "requirements.age_min", "requirements.gender_pref",
                           "requirements.experience_min", "requirements.education_required"):
            self.assertNotIn(khong_duoc, se_ghi)
        self.assertIn("requirements.japanese_required", se_ghi)

    def test_them_moi_van_dung_mac_dinh_cho_cot_vang(self):
        row = importer.parse_workbook(workbook_bytes([valid_row()], headers=self.BAT_BUOC[:-1])).rows[0]
        self.assertEqual(row.action, "create")
        self.assertIn("published", row.data)
        self.assertIn("status", row.data)
        self.assertEqual(row.data["requirements"]["gender_pref"], "khong_yeu_cau")
        self.assertNotIn("consultant_note", row.data)

    def test_cot_co_mat_o_trong_thi_van_xoa_co_chu_y(self):
        row = importer.parse_workbook(workbook_bytes([valid_row(code="DH-0001", salary_min=None, salary_max=None)])).rows[0]
        self.assertIn("salary_min", row.data["reference"])
        self.assertIsNone(row.data["reference"]["salary_min"])


class GioiHanDongTests(unittest.TestCase):
    """Quá 500 dòng thì từ chối cả file, nói rõ số dòng — trước 12/10 dòng 501+ bị bỏ im lặng."""

    def test_qua_gioi_han_thi_tu_choi_va_noi_so_dong(self):
        with self.assertRaises(ValueError) as loi:
            importer.parse_workbook(workbook_bytes([valid_row()] * (importer.MAX_ROWS + 1)))
        self.assertIn(f"hơn {importer.MAX_ROWS} dòng", str(loi.exception))
        self.assertIn(f"hàng {importer.MAX_ROWS + 2}", str(loi.exception))  # hàng 1 là tiêu đề
        self.assertIn("chia thành nhiều file", str(loi.exception))

    def test_dung_gioi_han_thi_nhan_du(self):
        preview = importer.parse_workbook(workbook_bytes([valid_row()] * importer.MAX_ROWS))
        self.assertEqual(preview.summary["total"], importer.MAX_ROWS)

    def test_dong_trong_khong_tinh_vao_gioi_han(self):
        preview = importer.parse_workbook(workbook_bytes([valid_row()] * importer.MAX_ROWS + [{}] * 20))
        self.assertEqual(preview.summary["total"], importer.MAX_ROWS)


class ImportEndpointTests(unittest.IsolatedAsyncioTestCase):
    """Chốt chặn ở tầng API, trước khi file kịp đi vào bộ đọc."""

    manager = {"email": "manager@example.com", "full_name": "Quản lý", "role": "manager"}

    def _request(self):
        from fastapi import Request

        return Request({"type": "http", "headers": [], "client": ("127.0.0.1", 1)})

    def _upload(self, filename: str, content: bytes):
        from fastapi import UploadFile

        return UploadFile(filename=filename, file=BytesIO(content))

    async def test_non_xlsx_file_is_refused(self):
        from fastapi import HTTPException

        from app.api.job_orders import import_job_orders

        with self.assertRaises(HTTPException) as raised:
            await import_job_orders(
                self._request(), self._upload("danhsach.csv", b"a,b"), True, self.manager
            )
        self.assertEqual(raised.exception.status_code, 415)

    async def test_oversized_file_is_refused_before_parsing(self):
        from fastapi import HTTPException

        from app.api.job_orders import MAX_IMPORT_BYTES, import_job_orders

        oversized = b"0" * (MAX_IMPORT_BYTES + 1)
        with self.assertRaises(HTTPException) as raised:
            await import_job_orders(
                self._request(), self._upload("to.xlsx", oversized), True, self.manager
            )
        self.assertEqual(raised.exception.status_code, 413)

    async def test_file_missing_required_columns_explains_what_to_do(self):
        from fastapi import HTTPException

        from app.api.job_orders import import_job_orders

        labels = [column.label for column in importer.COLUMNS if column.key != "title"]
        content = workbook_bytes([valid_row()], headers=labels)
        with self.assertRaises(HTTPException) as raised:
            await import_job_orders(
                self._request(), self._upload("thieu.xlsx", content), True, self.manager
            )
        self.assertEqual(raised.exception.status_code, 400)
        self.assertIn("Tên đơn", raised.exception.detail)


if __name__ == "__main__":
    unittest.main()
