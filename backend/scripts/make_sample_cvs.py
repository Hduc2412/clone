"""Sinh file CV mẫu để kiểm thử bộ đọc hồ sơ.

    python -m scripts.make_sample_cvs
    python -m scripts.make_sample_cvs --output duong/dan/khac

Tạo ra bốn định dạng khác nhau, vì bộ đọc phải xử lý được cả bốn:

- **PDF có lớp chữ**: đọc thẳng bằng PyMuPDF, nhanh và chính xác.
- **PDF bố cục bảng**: chữ nằm trong ô, thứ tự đọc dễ bị đảo. Đây là kiểu CV
  rất phổ biến ngoài đời và cũng là kiểu hay làm bộ đọc sai nhất.
- **DOCX**: đọc bằng python-docx.
- **PDF scan**: chỉ có ảnh, không có lớp chữ. Buộc phải nhận dạng chữ trong ảnh.
  Đây là trường hợp khó nhất và cũng rất hay gặp, vì nhiều người chụp ảnh CV
  giấy bằng điện thoại rồi gửi.

Kèm theo là `dap_an.json` ghi những gì bộ đọc **phải** rút ra đúng từ mỗi file.
Dùng để chấm điểm độ chính xác, không phải để làm đầu vào.
"""
import argparse
import json
from pathlib import Path
from typing import Any

import fitz  # PyMuPDF
from docx import Document
from docx.shared import Pt

# Chỉ để lấy tác dụng phụ: `app/__init__.py` chuyển đầu ra console sang UTF-8.
# Thiếu dòng này thì in tên file tiếng Việt ra màn hình Windows sẽ lỗi bảng mã,
# dù file vẫn được sinh đúng.
import app  # noqa: F401
from scripts.seed_data.sample_cvs import SAMPLE_CVS


BACKEND_DIR = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = BACKEND_DIR / "tests" / "fixtures" / "cv"

# Phông hệ thống có đủ dấu tiếng Việt. Phông dựng sẵn của PyMuPDF chỉ có bảng mã
# Latin-1 nên "ữ", "ặ", "ỡ" sẽ ra ô vuông — CV mẫu mà mất dấu thì không kiểm thử
# được gì, vì chính phần dấu mới là chỗ bộ đọc hay sai.
FONT_CANDIDATES = [
    Path("C:/Windows/Fonts/arial.ttf"),
    Path("C:/Windows/Fonts/segoeui.ttf"),
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
]

PAGE_WIDTH, PAGE_HEIGHT = fitz.paper_size("a4")
MARGIN = 56


def find_font() -> Path:
    for candidate in FONT_CANDIDATES:
        if candidate.exists():
            return candidate
    raise SystemExit(
        "Không tìm thấy phông chữ hỗ trợ tiếng Việt. "
        f"Đã thử: {', '.join(str(path) for path in FONT_CANDIDATES)}"
    )


class PdfWriter:
    """Ghi chữ tuần tự lên trang A4, tự sang trang khi hết chỗ."""

    def __init__(self, document: fitz.Document, font_path: Path) -> None:
        self.document = document
        self.font_path = str(font_path)
        self.page = document.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
        self.y = MARGIN

    def _ensure_space(self, needed: float) -> None:
        if self.y + needed > PAGE_HEIGHT - MARGIN:
            self.page = self.document.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
            self.y = MARGIN

    def text(
        self,
        value: str,
        *,
        size: float = 11,
        bold: bool = False,
        indent: float = 0,
        gap: float = 4,
        color: tuple[float, float, float] = (0.1, 0.1, 0.1),
    ) -> None:
        line_height = size * 1.45
        self._ensure_space(line_height + gap)
        self.page.insert_text(
            fitz.Point(MARGIN + indent, self.y + size),
            value,
            fontsize=size,
            fontfile=self.font_path,
            fontname="viet" + ("b" if bold else ""),
            color=color,
            render_mode=2 if bold else 0,  # tô viền để giả chữ đậm
        )
        self.y += line_height + gap

    def rule(self) -> None:
        self._ensure_space(10)
        self.page.draw_line(
            fitz.Point(MARGIN, self.y),
            fitz.Point(PAGE_WIDTH - MARGIN, self.y),
            color=(0.75, 0.75, 0.75),
            width=0.7,
        )
        self.y += 10

    def space(self, amount: float = 8) -> None:
        self.y += amount

    def table_row(self, left: str, right: str, size: float = 11) -> None:
        """Một dòng hai cột, dùng cho CV bố cục bảng."""
        line_height = size * 1.45
        self._ensure_space(line_height + 4)
        column = MARGIN + 150
        self.page.insert_text(
            fitz.Point(MARGIN + 6, self.y + size),
            left,
            fontsize=size,
            fontfile=self.font_path,
            fontname="viet",
            color=(0.35, 0.35, 0.35),
        )
        self.page.insert_text(
            fitz.Point(column + 6, self.y + size),
            right,
            fontsize=size,
            fontfile=self.font_path,
            fontname="viet",
            color=(0.1, 0.1, 0.1),
        )
        self.page.draw_rect(
            fitz.Rect(MARGIN, self.y - 2, PAGE_WIDTH - MARGIN, self.y + line_height),
            color=(0.8, 0.8, 0.8),
            width=0.6,
        )
        self.page.draw_line(
            fitz.Point(column, self.y - 2),
            fitz.Point(column, self.y + line_height),
            color=(0.8, 0.8, 0.8),
            width=0.6,
        )
        self.y += line_height + 4


def build_pdf(entry: dict[str, Any], font_path: Path) -> fitz.Document:
    content = entry["content"]
    document = fitz.open()
    writer = PdfWriter(document, font_path)

    writer.text(content["full_name"], size=19, bold=True, gap=2)
    writer.text(content["title"], size=12, color=(0.45, 0.45, 0.45))
    writer.rule()

    if entry["layout"] == "bang":
        for line in content["contact"]:
            label, _, value = line.partition(":")
            writer.table_row(label.strip(), value.strip() or label.strip())
    else:
        for line in content["contact"]:
            writer.text(line, size=10.5, color=(0.3, 0.3, 0.3), gap=1)

    writer.space(10)
    for heading, lines in content["sections"]:
        writer.text(heading, size=12, bold=True, gap=3)
        for line in lines:
            if entry["layout"] == "bang":
                writer.table_row("", line, size=10.5)
            else:
                writer.text(line, size=10.5, indent=10, gap=2)
        writer.space(8)

    return document


def build_scan(entry: dict[str, Any], font_path: Path) -> fitz.Document:
    """PDF chỉ chứa ảnh, không có lớp chữ.

    Dựng bằng cách in ra PDF chữ rồi chụp lại thành ảnh ở độ phân giải vừa phải
    và nhúng vào trang mới. Kết quả giống hệt tình huống người dùng chụp ảnh CV
    giấy rồi gửi: nhìn thì đọc được nhưng `get_text()` trả về rỗng.
    """
    source = build_pdf(entry, font_path)
    scanned = fitz.open()
    for page in source:
        # 130 dpi và nén JPEG: đúng tầm ảnh chụp bằng điện thoại, và giữ file ở
        # mức vài trăm KB thay vì vài MB. Ảnh quá nét cũng làm bài kiểm thử dễ
        # hơn thực tế, nên hạ độ phân giải là có chủ ý chứ không chỉ để nhẹ file.
        pixmap = page.get_pixmap(dpi=130)
        new_page = scanned.new_page(width=page.rect.width, height=page.rect.height)
        new_page.insert_image(new_page.rect, stream=pixmap.tobytes("jpeg", jpg_quality=72))
    source.close()
    return scanned


def build_docx(entry: dict[str, Any], path: Path) -> None:
    content = entry["content"]
    document = Document()

    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)

    heading = document.add_paragraph()
    run = heading.add_run(content["full_name"])
    run.bold = True
    run.font.size = Pt(18)

    subtitle = document.add_paragraph()
    subtitle.add_run(content["title"]).font.size = Pt(11)

    for line in content["contact"]:
        document.add_paragraph(line)

    for section_heading, lines in content["sections"]:
        paragraph = document.add_paragraph()
        section_run = paragraph.add_run(section_heading)
        section_run.bold = True
        section_run.font.size = Pt(12)
        for line in lines:
            document.add_paragraph(line, style="List Bullet")

    document.save(path)


def generate(output_dir: Path) -> dict[str, Any]:
    font_path = find_font()
    output_dir.mkdir(parents=True, exist_ok=True)

    report: list[dict[str, Any]] = []
    for entry in SAMPLE_CVS:
        slug = entry["slug"]
        kind = entry["format"]

        if kind == "docx":
            path = output_dir / f"{slug}.docx"
            build_docx(entry, path)
        elif kind == "pdf_scan":
            path = output_dir / f"{slug}_scan.pdf"
            document = build_scan(entry, font_path)
            document.save(path)
            document.close()
        else:
            path = output_dir / f"{slug}.pdf"
            document = build_pdf(entry, font_path)
            # Nhúng cả bộ phông Arial làm mỗi file nặng hơn một megabyte. Chỉ giữ
            # lại những chữ cái thực sự dùng là xuống còn vài chục kilobyte, mà
            # nội dung và dấu tiếng Việt vẫn nguyên.
            document.subset_fonts()
            document.save(path, garbage=4, deflate=True)
            document.close()

        report.append(
            {
                "file": path.name,
                "format": kind,
                "layout": entry["layout"],
                "note": entry["note"],
                "expected": entry["expected"],
                "size_bytes": path.stat().st_size,
            }
        )

    answers = output_dir / "dap_an.json"
    answers.write_text(
        json.dumps(report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    return {"files": report, "answers": str(answers)}


def verify(output_dir: Path) -> list[str]:
    """Kiểm tra lại các file vừa sinh: đọc được chữ chưa, bản scan có rỗng chữ không."""
    notes: list[str] = []
    for path in sorted(output_dir.glob("*.pdf")):
        with fitz.open(path) as document:
            text = "".join(page.get_text("text") for page in document)
        length = len(text.strip())
        if "_scan" in path.name:
            status = "ĐÚNG (không có lớp chữ, buộc phải nhận dạng ảnh)" if length < 50 else f"SAI: vẫn đọc được {length} ký tự"
        else:
            status = f"đọc được {length} ký tự" if length > 200 else f"SAI: chỉ {length} ký tự"
        notes.append(f"  {path.name}: {status}")

    for path in sorted(output_dir.glob("*.docx")):
        document = Document(path)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        notes.append(f"  {path.name}: đọc được {len(text.strip())} ký tự")
    return notes


def main() -> None:
    parser = argparse.ArgumentParser(description="Sinh CV mẫu để kiểm thử bộ đọc hồ sơ.")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    result = generate(args.output)
    print(f"Đã sinh {len(result['files'])} file CV mẫu tại {args.output}")
    for item in result["files"]:
        print(f"  {item['file']:<34} {item['format']:<9} {item['size_bytes']:>7} bytes  {item['note']}")
    print(f"\nĐáp án chấm điểm: {result['answers']}")
    print("\nKiểm tra lại:")
    for line in verify(args.output):
        print(line)


if __name__ == "__main__":
    main()
