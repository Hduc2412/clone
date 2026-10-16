"""Sinh bộ CV tiệm cận hồ sơ thật (`seed_data/cv_thuc_te.py`) vào `tests/fixtures/cv_thuc_te/`.

    python -m scripts.make_cv_thuc_te

Khác `make_sample_cvs.py` ở cách dựng: mỗi CV được **soạn thành tài liệu Word** —
bảng, ô tô màu, ảnh đại diện, đầu trang và chân trang — rồi **chính Microsoft Word
xuất ra PDF**, đúng đường ứng viên làm ngoài đời (soạn trong Word hoặc tải mẫu về
sửa, rồi "Lưu thành PDF"). Nhờ vậy lớp chữ trong PDF có đúng những tật của PDF
thật: thứ tự đọc của hai cột, chữ trong ô bảng, đầu trang lặp lại.

Cần Windows có Microsoft Word để xuất PDF. Các tệp kết quả được lưu vào kho mã nên
người khác không phải chạy lại. Không ghi đè tám CV ở `tests/fixtures/cv/`.
"""
import argparse
import json
import math
import random
import subprocess
import tempfile
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFilter

import app  # noqa: F401  — stdout UTF-8
from scripts.seed_data.cv_thuc_te import CV_THUC_TE

BACKEND_DIR = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = BACKEND_DIR / "tests" / "fixtures" / "cv_thuc_te"
XANH = RGBColor(0x1F, 0x4E, 0x79)


# ---------------------------------------------------------------- tiện ích Word
def nen_o(cell, mau: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), mau)
    tcpr.append(sh)


def khung_bang(table, co: bool = True) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for canh in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{canh}")
        b.set(qn("w:val"), "single" if co else "nil")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:color"), "808080")
        borders.append(b)
    tbl_pr.append(borders)


def phong(run, ten: str, co: float | None = None, dam: bool = False, mau: RGBColor | None = None,
          dong_a: str | None = None) -> None:
    run.font.name = ten
    rpr = run._element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), dong_a or ten)
    if co:
        run.font.size = Pt(co)
    run.bold = dam
    if mau is not None:
        run.font.color.rgb = mau


def doan(noi, chu: str, *, ten="Arial", co=10.5, dam=False, mau=None, canh=None, sau=2, dong_a=None):
    p = noi.add_paragraph()
    p.paragraph_format.space_after = Pt(sau)
    p.paragraph_format.space_before = Pt(0)
    if canh:
        p.alignment = canh
    phong(p.add_run(chu), ten, co, dam, mau, dong_a)
    return p


def anh_dai_dien(duong: Path) -> Path:
    """Hình bóng người xám trên nền nhạt — chỗ đặt ảnh thẻ, không phải ảnh người thật."""
    im = Image.new("RGB", (300, 400), (226, 232, 240))
    d = ImageDraw.Draw(im)
    d.ellipse((95, 70, 205, 190), fill=(150, 160, 175))
    d.ellipse((40, 215, 260, 470), fill=(150, 160, 175))
    im.save(duong)
    return duong


# ---------------------------------------------------------------- các kiểu trình bày
def kieu_chuan(doc, e, tieng_anh=False):
    doan(doc, e["ho_ten"], co=20, dam=True, mau=XANH, sau=0)
    if e.get("tieu_de"):
        doan(doc, e["tieu_de"], co=12, mau=RGBColor(0x59, 0x59, 0x59), sau=6)
    for nhan, gt in e["lien_he"]:
        doan(doc, f"{nhan}: {gt}", co=10.5, sau=0)
    for tieu_de, dong in e["muc"]:
        doan(doc, "", sau=0)
        doan(doc, tieu_de, co=12, dam=True, mau=XANH, sau=2)
        for d in dong:
            doan(doc, d, co=10.5, sau=1)


def kieu_hai_cot(doc, e, anh: Path):
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    khung_bang(t, co=False)
    trai, phai = t.rows[0].cells
    trai.width, phai.width = Cm(6.0), Cm(11.5)
    nen_o(trai, "DCE6F0")
    trai.paragraphs[0].add_run().add_picture(str(anh), width=Cm(3.2))
    trai.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doan(trai, "THÔNG TIN LIÊN HỆ", co=11, dam=True, mau=XANH, sau=2)
    for nhan, gt in e["lien_he"]:
        doan(trai, f"{nhan}: {gt}", co=9.5, sau=1)
    for tieu_de, dong in e.get("cot_trai", []):
        doan(trai, "", sau=0)
        doan(trai, tieu_de, co=11, dam=True, mau=XANH, sau=2)
        for d in dong:
            doan(trai, f"• {d}", co=9.5, sau=1)
    phai.paragraphs[0].add_run()
    doan(phai, e["ho_ten"], co=22, dam=True, mau=XANH, sau=0)
    doan(phai, e.get("tieu_de", ""), co=12, mau=RGBColor(0x59, 0x59, 0x59), sau=8)
    for tieu_de, dong in e["muc"]:
        doan(phai, tieu_de, co=12, dam=True, mau=XANH, sau=2)
        for d in dong:
            doan(phai, d, co=10.5, sau=1)
        doan(phai, "", sau=2)


def kieu_canva(doc, e):
    t = doc.add_table(rows=1, cols=1)
    khung_bang(t, co=False)
    o = t.rows[0].cells[0]
    nen_o(o, "2E4057")
    doan(o, e["ho_ten"], co=24, dam=True, mau=RGBColor(0xFF, 0xFF, 0xFF), canh=WD_ALIGN_PARAGRAPH.CENTER, sau=0)
    doan(o, e["tieu_de"].upper(), co=11, mau=RGBColor(0xF2, 0xC1, 0x4E), canh=WD_ALIGN_PARAGRAPH.CENTER, sau=4)
    doan(doc, "  |  ".join(f"{n}: {g}" for n, g in e["lien_he"]), co=9.5, canh=WD_ALIGN_PARAGRAPH.CENTER, sau=8)
    doan(doc, "KỸ NĂNG", co=12, dam=True, mau=RGBColor(0x2E, 0x40, 0x57), sau=2)
    for ten, diem in e["ky_nang"]:
        doan(doc, f"{ten:<16} {'●' * diem}{'○' * (5 - diem)}", co=10.5, sau=1)
    for tieu_de, dong in e["muc"]:
        doan(doc, "", sau=0)
        doan(doc, tieu_de, co=12, dam=True, mau=RGBColor(0x2E, 0x40, 0x57), sau=2)
        for d in dong:
            doan(doc, d, co=10.5, sau=1)


def kieu_so_yeu_ly_lich(doc, e, anh: Path):
    tnr = "Times New Roman"
    doan(doc, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", ten=tnr, co=12, dam=True, canh=WD_ALIGN_PARAGRAPH.CENTER, sau=0)
    doan(doc, "Độc lập – Tự do – Hạnh phúc", ten=tnr, co=12, dam=True, canh=WD_ALIGN_PARAGRAPH.CENTER, sau=0)
    doan(doc, "———————————", ten=tnr, co=10, canh=WD_ALIGN_PARAGRAPH.CENTER, sau=6)
    dau = doc.add_table(rows=1, cols=2)
    khung_bang(dau, co=False)
    o_anh, o_tieu = dau.rows[0].cells
    o_anh.width, o_tieu.width = Cm(3.5), Cm(13.5)
    o_anh.paragraphs[0].add_run().add_picture(str(anh), width=Cm(3.0))
    doan(o_anh, "Ảnh 4x6", ten=tnr, co=9, canh=WD_ALIGN_PARAGRAPH.CENTER)
    doan(o_tieu, "SƠ YẾU LÝ LỊCH TỰ THUẬT", ten=tnr, co=18, dam=True, canh=WD_ALIGN_PARAGRAPH.CENTER, sau=6)
    for nhan, gt in e["truong"]:
        doan(doc, f"{nhan}: {gt}", ten=tnr, co=12, sau=2)

    def bang(tieu_de, cot, hang):
        doan(doc, tieu_de, ten=tnr, co=12, dam=True, sau=2)
        t = doc.add_table(rows=1, cols=len(cot))
        khung_bang(t)
        for i, c in enumerate(cot):
            o = t.rows[0].cells[i]
            o.text = ""
            phong(o.paragraphs[0].add_run(c), tnr, 11, True)
        for h in hang or [("",) * len(cot)]:
            cs = t.add_row().cells
            for i, v in enumerate(h):
                cs[i].text = ""
                phong(cs[i].paragraphs[0].add_run(v), tnr, 11)
        doan(doc, "", sau=2)

    bang("I. QUAN HỆ GIA ĐÌNH", ["Quan hệ", "Họ và tên", "Năm sinh", "Nghề nghiệp", "Điện thoại"], e["gia_dinh"])
    bang("II. QUÁ TRÌNH ĐÀO TẠO", ["Thời gian", "Tên trường", "Ngành học", "Văn bằng"], e["dao_tao"])
    bang("III. QUÁ TRÌNH CÔNG TÁC", ["Thời gian", "Đơn vị", "Công việc"], e["cong_tac"])
    doan(doc, "Tôi xin cam đoan những lời khai trên là đúng sự thật. Nếu có gì sai, tôi xin hoàn toàn "
              "chịu trách nhiệm.", ten=tnr, co=12, sau=6)
    doan(doc, "Ngày ...... tháng ...... năm 2026", ten=tnr, co=12, canh=WD_ALIGN_PARAGRAPH.RIGHT, sau=0)
    doan(doc, "Người khai", ten=tnr, co=12, dam=True, canh=WD_ALIGN_PARAGRAPH.RIGHT, sau=24)
    doan(doc, e["ho_ten"].title(), ten=tnr, co=12, canh=WD_ALIGN_PARAGRAPH.RIGHT)


def kieu_rirekisho(doc, e):
    r = e["rirekisho"]
    jp = "MS Mincho"

    def o_chu(cell, chu, co=10.5, dam=False, canh=None):
        cell.text = ""
        p = cell.paragraphs[0]
        if canh:
            p.alignment = canh
        phong(p.add_run(chu), jp, co, dam, dong_a=jp)

    doan(doc, "履　歴　書", ten=jp, co=20, dam=True, sau=0, dong_a=jp)
    doan(doc, r["ngay"], ten=jp, co=10, canh=WD_ALIGN_PARAGRAPH.RIGHT, sau=4, dong_a=jp)
    t = doc.add_table(rows=5, cols=2)
    khung_bang(t)
    for i, (nhan, gt) in enumerate([("ふりがな", r["furigana"]), ("氏名", r["ten"]),
                                    ("生年月日", f"{r['sinh']}　　性別　{r['gioi']}"),
                                    ("現住所", r["dia_chi"]), ("電話", r["dien_thoai"])]):
        o_chu(t.rows[i].cells[0], nhan, 9)
        o_chu(t.rows[i].cells[1], gt, 16 if nhan == "氏名" else 10.5, dam=nhan == "氏名")
        t.rows[i].cells[0].width, t.rows[i].cells[1].width = Cm(2.5), Cm(14.5)
    doan(doc, "", sau=4)

    def bang_nam_thang(tieu, hang):
        t = doc.add_table(rows=1, cols=3)
        khung_bang(t)
        for i, c in enumerate(["年", "月", tieu]):
            o_chu(t.rows[0].cells[i], c, 10, True, WD_ALIGN_PARAGRAPH.CENTER)
        for nam, thang, nd in hang:
            cs = t.add_row().cells
            o_chu(cs[0], nam)
            o_chu(cs[1], thang)
            o_chu(cs[2], nd, canh=WD_ALIGN_PARAGRAPH.CENTER if nd in ("学歴", "職歴") else
                  (WD_ALIGN_PARAGRAPH.RIGHT if nd == "以上" else None))
        for row in t.rows:
            row.cells[0].width, row.cells[1].width, row.cells[2].width = Cm(1.8), Cm(1.2), Cm(14.0)
        doan(doc, "", sau=4)

    bang_nam_thang("学歴・職歴", r["hoc_lam"])
    bang_nam_thang("免許・資格", r["chung_chi"])
    t = doc.add_table(rows=2, cols=1)
    khung_bang(t)
    o_chu(t.rows[0].cells[0], "志望の動機・自己PR\n" + r["dong_co"])
    o_chu(t.rows[1].cells[0], "本人希望記入欄\n" + r["nguyen_vong"])


def kieu_bang_docx(doc, e, anh: Path):
    doan(doc, "SƠ YẾU HỒ SƠ ỨNG VIÊN", co=16, dam=True, mau=XANH, canh=WD_ALIGN_PARAGRAPH.CENTER, sau=6)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(anh), width=Cm(2.6))
    t = doc.add_table(rows=0, cols=2)
    khung_bang(t)
    for nhan, gt in e["bang"]:
        cs = t.add_row().cells
        nen_o(cs[0], "EEF2F7")
        cs[0].text, cs[1].text = "", ""
        phong(cs[0].paragraphs[0].add_run(nhan), "Arial", 10.5, True)
        phong(cs[1].paragraphs[0].add_run(gt), "Arial", 10.5)
        cs[0].width, cs[1].width = Cm(4.0), Cm(13.0)


def kieu_mau_trung_tam(doc, e, anh: Path):
    doan(doc, "PHIẾU THÔNG TIN ỨNG VIÊN / 応募者情報", co=15, dam=True, mau=XANH,
         canh=WD_ALIGN_PARAGRAPH.CENTER, sau=2, dong_a="MS Gothic")
    doan(doc, "Chương trình Kỹ năng đặc định – ngành Điều dưỡng (介護)", co=10, canh=WD_ALIGN_PARAGRAPH.CENTER,
         sau=6, dong_a="MS Gothic")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run().add_picture(str(anh), width=Cm(2.6))
    t = doc.add_table(rows=0, cols=2)
    khung_bang(t)
    for nhan, gt in e["phieu"]:
        cs = t.add_row().cells
        nen_o(cs[0], "F2F2F2")
        cs[0].text, cs[1].text = "", ""
        phong(cs[0].paragraphs[0].add_run(nhan), "Arial", 10, True, dong_a="MS Gothic")
        phong(cs[1].paragraphs[0].add_run(gt), "Arial", 10.5, dong_a="MS Gothic")
        cs[0].width, cs[1].width = Cm(5.0), Cm(12.0)


def kieu_khong_dau(doc, e):
    for d in e["dong"]:
        doan(doc, d, ten="Calibri", co=11, sau=3)


def kieu_hai_trang(doc, e):
    doan(doc, "THƯ XIN VIỆC", co=16, dam=True, canh=WD_ALIGN_PARAGRAPH.CENTER, sau=12)
    for d in e["thu"]:
        doan(doc, d, co=11, sau=8)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    kieu_chuan(doc, e)


def kieu_header_footer(doc, e):
    sec = doc.sections[0]
    phong(sec.header.paragraphs[0].add_run(f"CV – {e['ho_ten'].title()}"), "Arial", 8, mau=RGBColor(0x80, 0x80, 0x80))
    sec.header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    phong(p.add_run("Trang 1/1 – Hồ sơ ứng tuyển chương trình điều dưỡng Nhật Bản"), "Arial", 8,
          mau=RGBColor(0x80, 0x80, 0x80))
    kieu_chuan(doc, e)


def dung_docx(e: dict, duong: Path, anh: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.6)
    kieu = e["layout"]
    if kieu == "chuan":
        kieu_chuan(doc, e)
    elif kieu == "hai_cot":
        kieu_hai_cot(doc, e, anh)
    elif kieu == "canva":
        kieu_canva(doc, e)
    elif kieu == "so_yeu_ly_lich":
        kieu_so_yeu_ly_lich(doc, e, anh)
    elif kieu == "rirekisho":
        kieu_rirekisho(doc, e)
    elif kieu == "bang_docx":
        kieu_bang_docx(doc, e, anh)
    elif kieu == "mau_trung_tam":
        kieu_mau_trung_tam(doc, e, anh)
    elif kieu == "khong_dau":
        kieu_khong_dau(doc, e)
    elif kieu == "hai_trang":
        kieu_hai_trang(doc, e)
    elif kieu == "header_footer":
        kieu_header_footer(doc, e)
    else:
        raise ValueError(kieu)
    doc.save(duong)


# ---------------------------------------------------------------- xuất PDF và ảnh
WORD_PS = r"""
param([string]$Danh)
$w = New-Object -ComObject Word.Application
$w.Visible = $false; $w.DisplayAlerts = 0
try {
  Get-Content -Encoding UTF8 $Danh | ForEach-Object {
    $vao, $ra = $_ -split "`t"
    if (Test-Path $ra) { Remove-Item $ra -Force }
    $d = $w.Documents.Open($vao, $false, $true)
    $d.ExportAsFixedFormat($ra, 17)
    $d.Close($false)
    if (-not (Test-Path $ra)) { throw "Khong xuat duoc $ra" }
  }
} finally { $w.Quit() }
"""


def word_ra_pdf(cap: list[tuple[Path, Path]], tam: Path) -> None:
    ps = tam / "xuat.ps1"
    ps.write_text(WORD_PS, encoding="utf-8")
    danh = tam / "danh.txt"
    danh.write_text("\n".join(f"{a}\t{b}" for a, b in cap), encoding="utf-8")
    subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(ps), str(danh)],
                   check=True)


def anh_chup_dien_thoai(pdf: Path, ra: Path) -> None:
    """Như chụp tờ giấy bằng điện thoại: hơi nghiêng, một góc tối, nén JPEG."""
    with fitz.open(pdf) as d:
        pix = d[0].get_pixmap(dpi=110)
    im = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    nen = Image.new("RGB", (im.width + 160, im.height + 160), (92, 84, 74))
    nen.paste(im, (80, 80))
    nen = nen.rotate(2.3, resample=Image.BICUBIC, fillcolor=(92, 84, 74))
    bong = Image.new("L", nen.size, 0)
    ImageDraw.Draw(bong).rectangle((0, 0, nen.width, nen.height), fill=0)
    w, h = nen.size
    px = bong.load()
    for y in range(0, h, 4):
        for x in range(0, w, 4):
            v = int(70 * max(0.0, 1 - math.hypot(x - w, y) / (0.9 * w)))
            for dy in range(4):
                for dx in range(4):
                    if x + dx < w and y + dy < h:
                        px[x + dx, y + dy] = v
    toi = Image.new("RGB", nen.size, (0, 0, 0))
    nen = Image.composite(toi, nen, bong.filter(ImageFilter.GaussianBlur(25)))
    random.seed(14)
    nen.save(ra, "JPEG", quality=68)


def anh_scan(pdf: Path, ra: Path) -> None:
    with fitz.open(pdf) as src, fitz.open() as dst:
        for trang in src:
            pix = trang.get_pixmap(dpi=130)
            moi = dst.new_page(width=trang.rect.width, height=trang.rect.height)
            moi.insert_image(moi.rect, stream=pix.tobytes("jpeg", jpg_quality=72))
        dst.save(ra)


# ---------------------------------------------------------------- chạy
def generate(ra: Path) -> list[dict]:
    ra.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as t:
        tam = Path(t)
        anh = anh_dai_dien(tam / "anh_dai_dien.png")
        can_pdf, bao_cao = [], []
        for e in CV_THUC_TE:
            docx = tam / f"{e['slug']}.docx"
            dung_docx(e, docx, anh)
            if e["format"] == "docx":
                dich = ra / f"{e['slug']}.docx"
                dich.write_bytes(docx.read_bytes())
            else:
                can_pdf.append((docx, tam / f"{e['slug']}.pdf"))
        word_ra_pdf(can_pdf, tam)
        for e in CV_THUC_TE:
            pdf_tam = tam / f"{e['slug']}.pdf"
            if e["format"] == "pdf":
                dich = ra / f"{e['slug']}.pdf"
                dich.write_bytes(pdf_tam.read_bytes())
            elif e["format"] == "jpg":
                dich = ra / f"{e['slug']}.jpg"
                anh_chup_dien_thoai(pdf_tam, dich)
            elif e["format"] == "pdf_scan":
                dich = ra / f"{e['slug']}_scan.pdf"
                anh_scan(pdf_tam, dich)
            else:
                dich = ra / f"{e['slug']}.docx"
            bao_cao.append({"file": dich.name, "format": e["format"], "layout": e["layout"], "note": e["note"],
                            "tinh_huong": e["tinh_huong"], "expected": e["expected"],
                            "size_bytes": dich.stat().st_size})
    (ra / "dap_an.json").write_text(json.dumps(bao_cao, ensure_ascii=False, indent=2), encoding="utf-8")
    return bao_cao


def main() -> None:
    ap = argparse.ArgumentParser(description="Sinh bộ CV tiệm cận hồ sơ thật.")
    ap.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    a = ap.parse_args()
    ds = generate(a.output)
    print(f"Đã sinh {len(ds)} CV tại {a.output}")
    for x in ds:
        print(f"  {x['file']:<44} {x['format']:<5} {x['layout']:<15} {x['size_bytes']:>8} bytes")


if __name__ == "__main__":
    main()
