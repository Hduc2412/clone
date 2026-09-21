"""Đọc chữ ra khỏi file CV.

Rút chữ được bằng thư viện với PDF và DOCX. Hai định dạng này đọc được **không
cần gọi mô hình ngôn ngữ**, nên đó là đường chính.

## Vì sao vẫn nhận ảnh dù chưa đọc được

Ảnh chụp bằng điện thoại thì chưa rút được chữ — muốn đọc phải nhờ mô hình, mà
việc đó chưa làm. Nhưng **vẫn nhận và cất giữ**, thay vì chặn ở cửa bằng một lỗi.

Lý do là để trả lời một câu hỏi mà hiện không ai biết đáp án: có bao nhiêu ứng
viên thật sự gửi ảnh chụp thay vì gửi file? Chặn ngay ở cửa thì hệ thống không
bao giờ học được điều đó — người dùng gặp lỗi rồi bỏ đi, không để lại dấu vết
nào. Nhận vào thì mỗi lần như vậy thành một bản ghi có trạng thái "chưa đọc được
chữ", và sau vài tuần chạy thật là có con số thật để quyết định có đáng làm phần
đọc ảnh hay không.

Chi phí của việc nhận vào gần bằng không: không gọi mô hình, chỉ ghi một bản ghi
và cất file. Đổi lại, ứng viên nhận được lời mời khai tay thay vì một thông báo
lỗi cụt lủn, và nhân viên vẫn mở được ảnh ra đọc.

## Vì sao kiểm tra cả phần mở đầu của file

Trình duyệt khai `content_type` theo đuôi file, mà đuôi file thì ai cũng đổi được.
Một file `.exe` đổi tên thành `.pdf` vẫn được khai là `application/pdf`. Nên ở đây
đọc vài byte đầu để xác nhận đúng định dạng thật: PDF bắt đầu bằng `%PDF-`, DOCX
là file ZIP nên bắt đầu bằng `PK`. Không thay thế được việc quét virus, nhưng
chặn được nhầm lẫn thông thường và phần lớn trò đổi đuôi.

## "Cần đọc bằng mô hình" nghĩa là gì

CV scan là ảnh nằm trong khung PDF: mở ra thấy chữ, nhưng rút chữ thì được chuỗi
rỗng. Trường hợp đó `needs_model_ocr` bật lên, và tầng trên quyết định có gửi bản
gốc cho mô hình đọc hay không — giới hạn số trang theo `cv_ocr_max_pages` để một
hồ sơ dày không đốt hết hạn mức gọi mô hình.
"""
from dataclasses import dataclass

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

JPEG_MIME = "image/jpeg"
PNG_MIME = "image/png"
WEBP_MIME = "image/webp"
HEIC_MIME = "image/heic"

# Ảnh chụp: **nhận và cất giữ, chưa đọc chữ**. Xem docstring mục "Vì sao vẫn nhận
# ảnh" ở đầu file.
IMAGE_TYPES: dict[str, str] = {
    JPEG_MIME: ".jpg",
    PNG_MIME: ".png",
    WEBP_MIME: ".webp",
    HEIC_MIME: ".heic",
}

SUPPORTED_TYPES: dict[str, str] = {
    PDF_MIME: ".pdf",
    DOCX_MIME: ".docx",
    **IMAGE_TYPES,
}

SUPPORTED_LABEL = "PDF, DOCX hoặc ảnh chụp"

# Số ký tự tối thiểu để coi là đã rút được chữ. Một trang PDF scan thường vẫn trả
# về vài ký tự rác từ phần chú thích hoặc watermark, nên ngưỡng 0 là quá lỏng.
MIN_MEANINGFUL_CHARS = 80

# Cắt bớt trước khi đưa cho mô hình. CV dài nhất trong thực tế hiếm khi quá chừng
# này, và để nguyên một file nhồi hàng trăm nghìn ký tự đi thẳng vào prompt là mở
# đường cho cả lỗi tràn lẫn chi phí không kiểm soát.
MAX_TEXT_CHARS = 20_000


class UnsupportedDocument(Exception):
    """File không thuộc định dạng nhận được, hoặc nội dung không khớp với đuôi."""


@dataclass(frozen=True)
class ReadResult:
    text: str
    page_count: int
    needs_model_ocr: bool


def detect_kind(data: bytes, declared_type: str, filename: str) -> str:
    """Trả về mime chuẩn của file, dựa trên nội dung thật chứ không tin phần khai.

    `declared_type` và đuôi file chỉ dùng để báo lỗi cho dễ hiểu; thứ quyết định
    là mấy byte đầu.
    """
    if data.startswith(b"%PDF-"):
        return PDF_MIME
    # DOCX là một file ZIP. `PK\x03\x04` là chữ ký của ZIP; các biến thể còn lại
    # (`PK\x05\x06` rỗng, `PK\x07\x08` chia nhỏ) không phải DOCX hợp lệ.
    if data.startswith(b"PK\x03\x04"):
        return DOCX_MIME

    image_kind = _detect_image(data)
    if image_kind is not None:
        return image_kind

    hint = filename or declared_type or "file"
    raise UnsupportedDocument(
        f"Chỉ nhận {SUPPORTED_LABEL}. Nội dung của {hint} không khớp định dạng nào "
        "trong số đó."
    )


def _detect_image(data: bytes) -> str | None:
    """Nhận diện ảnh theo mấy byte đầu. Trả `None` nếu không phải ảnh."""
    if data.startswith(b"\xff\xd8\xff"):
        return JPEG_MIME
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return PNG_MIME
    # WEBP là khung RIFF: 4 byte "RIFF", 4 byte độ dài, rồi 4 byte "WEBP".
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return WEBP_MIME
    # HEIC là khung ISO-BMFF: 4 byte độ dài, rồi "ftyp" và mã thương hiệu. Máy
    # iPhone mặc định chụp ra định dạng này, nên bỏ qua nó là bỏ qua một nửa số
    # ảnh người dùng gửi lên.
    if data[4:8] == b"ftyp" and data[8:12] in (b"heic", b"heix", b"heif", b"mif1", b"msf1"):
        return HEIC_MIME
    return None


def read_image(data: bytes) -> ReadResult:
    """Ảnh thì chưa rút được chữ — luôn phải nhờ mô hình đọc.

    Không gọi mô hình ở đây. Hàm này cố tình chỉ trả về "cần đọc bằng mô hình" để
    tầng trên ghi nhận file rồi mời ứng viên khai tay. Xem lý do ở đầu file.
    """
    return ReadResult(text="", page_count=1, needs_model_ocr=True)


def read_pdf(data: bytes) -> ReadResult:
    import pymupdf

    with pymupdf.open(stream=data, filetype="pdf") as document:
        pages = [page.get_text() for page in document]
        page_count = document.page_count

    text = _clean("\n".join(pages))
    return ReadResult(
        text=text,
        page_count=page_count,
        needs_model_ocr=len(text) < MIN_MEANINGFUL_CHARS,
    )


def read_docx(data: bytes) -> ReadResult:
    import io

    from docx import Document

    document = Document(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]

    # Rất nhiều CV trình bày bằng bảng. Bỏ qua bảng là mất đúng phần thông tin
    # đậm đặc nhất: học vấn, kinh nghiệm, trình độ tiếng.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    text = _clean("\n".join(parts))
    return ReadResult(
        text=text,
        page_count=0,  # DOCX không có khái niệm số trang cố định.
        needs_model_ocr=len(text) < MIN_MEANINGFUL_CHARS,
    )


def read(data: bytes, declared_type: str = "", filename: str = "") -> tuple[str, ReadResult]:
    """Đọc file, trả về mime chuẩn kèm kết quả."""
    kind = detect_kind(data, declared_type, filename)
    if kind in IMAGE_TYPES:
        return kind, read_image(data)
    result = read_pdf(data) if kind == PDF_MIME else read_docx(data)
    return kind, result


def _clean(text: str) -> str:
    """Bỏ khoảng trắng thừa, gộp dòng trống liên tiếp, cắt theo hạn mức."""
    lines = [line.strip() for line in text.replace("\r\n", "\n").split("\n")]
    cleaned: list[str] = []
    for line in lines:
        if line:
            cleaned.append(line)
        elif cleaned and cleaned[-1]:
            cleaned.append("")
    while cleaned and not cleaned[-1]:
        cleaned.pop()
    return "\n".join(cleaned)[:MAX_TEXT_CHARS]
